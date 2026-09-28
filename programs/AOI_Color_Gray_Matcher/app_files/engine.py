# -*- coding: utf-8 -*-
from pathlib import Path
import os,re,csv
# Pillow/openpyxl are imported inside the functions that use them, so program start
# (folder scan and UI) does not pay for loading them.
RESULT_FOLDER='AOI_Color_Gray_Matching_Result';MAX_WAFERS=25;TW,TH=360,270;GW,GH,GP=3168,1024,0.769644115336745
def transform_file(p):
 files=sorted(p.glob('frameToChuckPlane.*.ini'));return files[0] if files else None
def is_wafer(p):return (p/'ColorImageGrabingInfo.ini').exists() and (p/'ScanResultImageList.txt').exists() and transform_file(p) is not None
def child_dirs(p):
 try:
  with os.scandir(p) as entries:return sorted([Path(e.path) for e in entries if e.is_dir(follow_symlinks=False)],key=lambda x:x.name.lower())
 except (OSError,PermissionError):return []
def discover(root,callback=None,cancel=None,max_depth=5):
 root=Path(root)
 if is_wafer(root):
  callback and callback('found',root,1,1,0);return [root]
 skip={RESULT_FOLDER.lower(),'_crop_images','_excel_thumbnails','_metadata','$recycle.bin','system volume information'};queue=[(root,0)];seen=set();found=[];scanned=0
 while queue and len(found)<MAX_WAFERS:
  if cancel and cancel.is_set():break
  current,depth=queue.pop(0);key=str(current).lower()
  if key in seen:continue
  seen.add(key);scanned+=1;callback and callback('scanning',current,scanned,len(found),len(queue))
  if depth>=max_depth:continue
  for child in child_dirs(current):
   if child.name.lower() in skip:continue
   if is_wafer(child):found.append(child);callback and callback('found',child,scanned,len(found),len(queue))
   else:queue.append((child,depth+1))
 callback and callback('complete',root,scanned,len(found),len(queue));return found
def parse_colors(path):
 rows=[];current=None
 for raw in path.read_text(encoding='utf-8-sig',errors='replace').splitlines():
  line=raw.strip()
  if line.startswith('[') and line.endswith(']'):
   if current:rows.append(current)
   current={'file':line[1:-1]}
  elif current and '=' in line:
   key,value=line.split('=',1);current[key.strip()]=value.strip()
 if current:rows.append(current)
 return rows
def parse_frames(path):
 rows=[]
 for raw in path.read_text(encoding='utf-8-sig',errors='replace').splitlines():
  parts=raw.strip().split(',')
  if len(parts)<3 or parts[0].startswith('Version='):continue
  try:
   match=re.search(r'\.t\.(\d+)\.jpe?g$',parts[0],re.I);rows.append({'file':parts[0],'x':float(parts[1]),'y':float(parts[2]),'recipe':int(match.group(1)) if match else None})
  except ValueError:pass
 return rows
def parse_transform(path):
 data={}
 for raw in path.read_text(encoding='utf-8-sig',errors='replace').splitlines():
  if '=' in raw:
   key,value=raw.split('=',1)
   try:data[key.strip()]=[float(x) for x in value.split()]
   except ValueError:pass
 sx,sy=data['Scan_X'],data['Scan_Y'];matrix=((sx[0],sx[1]),(sy[0],sy[1]));offset=(sx[2],sy[2]);det=matrix[0][0]*matrix[1][1]-matrix[0][1]*matrix[1][0]
 return ((matrix[1][1]/det,-matrix[0][1]/det),(-matrix[1][0]/det,matrix[0][0]/det)),offset
def choose_frame(color,frames,inverse,offset):
 fx=float(color.get('FaultX',color.get('X')));fy=float(color.get('FaultY',color.get('Y')));recipe=int(color.get('RecipeNumber','1'));candidates=[]
 for frame in frames:
  if frame['recipe']!=recipe:continue
  tx=fx-frame['x']-offset[0];ty=fy-frame['y']-offset[1];px=inverse[0][0]*tx+inverse[0][1]*ty;py=inverse[1][0]*tx+inverse[1][1]*ty
  if 0<=px<GW and 0<=py<GH:candidates.append((min(px,GW-px,py,GH-py),frame,px,py))
 if not candidates:return None,None,None,0
 candidates.sort(key=lambda x:x[0],reverse=True);_,frame,px,py=candidates[0];return frame,px,py,len(candidates)
def save_thumb(source,destination,marker=None):
 from PIL import Image,ImageDraw,ImageOps
 with Image.open(source) as image:
  image=ImageOps.exif_transpose(image).convert('RGB');image.thumbnail((TW,TH),Image.Resampling.LANCZOS);canvas=Image.new('RGB',(TW,TH),'white');ox=(TW-image.width)//2;oy=(TH-image.height)//2;canvas.paste(image,(ox,oy))
  if marker:
   px,py,w,h=marker;x=ox+px*image.width/w;y=oy+py*image.height/h;draw=ImageDraw.Draw(canvas);draw.line((x-15,y,x+15,y),fill='red',width=3);draw.line((x,y-15,x,y+15),fill='red',width=3)
  canvas.save(destination,'JPEG',quality=78,optimize=True)
def excel_link(base,target):
 base=Path(base).resolve();target=Path(target).resolve()
 try:return '.\\'+str(target.relative_to(base)).replace('/','\\')
 except ValueError:return str(target)
def crop_gray(gray_path,crop_path,px,py,crop_w,crop_h,out_w,out_h):
 from PIL import Image,ImageOps
 with Image.open(gray_path) as image:
  image=ImageOps.exif_transpose(image).convert('L');aw,ah=image.size;x=px*aw/GW;y=py*ah/GH;left=round(x-crop_w/2);top=round(y-crop_h/2);canvas=Image.new('L',(crop_w,crop_h),0);sl=max(0,left);st=max(0,top);sr=min(aw,left+crop_w);sb=min(ah,top+crop_h);canvas.paste(image.crop((sl,st,sr,sb)),(sl-left,st-top));canvas.resize((out_w,out_h),Image.Resampling.LANCZOS).save(crop_path,'JPEG',quality=95);return(left,top,left+crop_w,top+crop_h)
def build_workbook(wafer,result,records,callback):
 from openpyxl import Workbook
 from openpyxl.drawing.image import Image as XLImage
 from openpyxl.styles import Alignment,Font,PatternFill
 output=result/f'{wafer.name}_Color_Gray_Crop.xlsx';wb=Workbook();ws=wb.active;ws.title='Image_Comparison';ws.sheet_view.showGridLines=False;ws.freeze_panes='A2';ws.append(['Color image','Gray image','Cropped gray image','Coordinates / matching details'])
 for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='111827');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(horizontal='center')
 for col in 'ABC':ws.column_dimensions[col].width=52
 ws.column_dimensions['D'].width=58;total=max(1,len(records));callback(0,total,'Excel 준비')
 for index,record in enumerate(records,1):
  row=index+1;ws.row_dimensions[row].height=205
  for col,thumb_key,path_key,label in zip('ABC',('color_thumb','gray_thumb','crop_thumb'),('color_path','gray_path','crop_path'),('Color 사진 열기','Gray 사진 열기','Crop 사진 열기')):
   cell=ws[f'{col}{row}'];cell.value=label;cell.hyperlink=excel_link(output.parent,record[path_key]);cell.style='Hyperlink';cell.alignment=Alignment(horizontal='center',vertical='bottom');picture=XLImage(str(record[thumb_key]));picture.width=TW;picture.height=245;ws.add_image(picture,f'{col}{row}')
  ws.cell(row,4,f"Color: {record['color_file']}\nGray: {record['gray_file']}\nFault: {record['fault_x']:.3f}, {record['fault_y']:.3f}\nGray Pixel: {record['pixel_x']:.1f}, {record['pixel_y']:.1f}\nCandidates: {record['candidate_count']}\nCrop: {record['crop_box']}");ws.cell(row,4).alignment=Alignment(wrap_text=True,vertical='top');callback(index,total,f'Excel 이미지 삽입 {index}/{total}')
 callback(total,total,'Excel 파일 저장');wb.save(output);callback(total,total,'Excel 생성 완료');return output
def process_wafer(wafer,output_root,progress,log,cancel,wafer_index,wafer_total):
 result=output_root/wafer.name;crops=result/'_crop_images';thumbs=result/'_excel_thumbnails';crops.mkdir(parents=True,exist_ok=True);thumbs.mkdir(parents=True,exist_ok=True);colors=parse_colors(wafer/'ColorImageGrabingInfo.ini');frames=parse_frames(wafer/'ScanResultImageList.txt');inverse,offset=parse_transform(transform_file(wafer));records=[];failures=[]
 for index,color in enumerate(colors,1):
  if cancel.is_set():raise RuntimeError('사용자가 작업을 취소했습니다.')
  progress(wafer_index,wafer_total,index-1,len(colors),wafer.name,color['file'],'이미지 매칭 및 Crop','image',index/max(1,len(colors)));frame,px,py,count=choose_frame(color,frames,inverse,offset);color_path=wafer/color['file']
  if not frame or not color_path.exists() or not (wafer/frame['file']).exists():failures.append((color['file'],'매칭 또는 파일 없음'));continue
  gray_path=wafer/frame['file'];out_w=int(float(color.get('ImageSizeX',1380)));out_h=int(float(color.get('ImageSizeY',1036)));crop_w=max(1,round(out_w*float(color.get('PixelSizeX',.4674))/GP));crop_h=max(1,round(out_h*float(color.get('PixelSizeY',.4674))/GP));crop_path=crops/f"{Path(color['file']).stem}_gray_crop.jpeg";crop_box=crop_gray(gray_path,crop_path,px,py,crop_w,crop_h,out_w,out_h)
  color_thumb=thumbs/f'{index:04d}_color.jpg';gray_thumb=thumbs/f'{index:04d}_gray.jpg';crop_thumb=thumbs/f'{index:04d}_crop.jpg';save_thumb(color_path,color_thumb);save_thumb(gray_path,gray_thumb,(px,py,GW,GH));save_thumb(crop_path,crop_thumb);records.append({'color_thumb':color_thumb,'gray_thumb':gray_thumb,'crop_thumb':crop_thumb,'color_path':color_path,'gray_path':gray_path,'crop_path':crop_path,'color_file':color['file'],'gray_file':frame['file'],'fault_x':float(color.get('FaultX',color.get('X'))),'fault_y':float(color.get('FaultY',color.get('Y'))),'pixel_x':px,'pixel_y':py,'candidate_count':count,'crop_box':crop_box})
 def excel_progress(done,total,message):progress(wafer_index,wafer_total,done,total,wafer.name,message,message,'excel',done/max(1,total));log(message)
 workbook=build_workbook(wafer,result,records,excel_progress)
 if failures:
  with (result/'unmatched_or_missing.csv').open('w',newline='',encoding='utf-8-sig') as handle:csv.writer(handle).writerows([('ColorFile','Reason'),*failures])
 return {'wafer':wafer.name,'workbook':str(workbook),'matched':len(records),'failed':len(failures)}
