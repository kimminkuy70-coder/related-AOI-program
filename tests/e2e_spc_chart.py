# -*- coding: utf-8 -*-
"""SPC Chart HTML (programs/SPC_Chart_HTML): the built single .html in Chromium (Edge's engine).

Edge's file dialog cannot be driven, so showOpenFilePicker hands out a file of the page's private
file system (OPFS): a real file handle, re-read the same way as a picked Excel file. Rewriting that
file stands in for saving the workbook in Excel. Each test runs in a fresh browser context.

    python -m unittest tests.e2e_spc_chart -v
"""
import base64
import re
import datetime
import functools
import http.server
import io
import json
import subprocess
import sys
import threading
import unittest
import zipfile
from pathlib import Path

try:
    from playwright.sync_api import sync_playwright
    import openpyxl
except ImportError:  # pragma: no cover
    sync_playwright = None

REPO = Path(__file__).resolve().parent.parent
PROGRAM = REPO / 'programs' / 'SPC_Chart_HTML'
VERSION = json.loads((PROGRAM / 'version.json').read_text(encoding='utf-8'))
PAGE = VERSION['release'] + '.html'

PICKER = """
window.showOpenFilePicker = async () => {
  const root = await navigator.storage.getDirectory();
  return [await root.getFileHandle(window.__pick, {create: true})];
};
"""

PUT_FILE = """async ([name, b64]) => {
  const root = await navigator.storage.getDirectory();
  const h = await root.getFileHandle(name, {create: true});
  const w = await h.createWritable();
  await w.write(Uint8Array.from(atob(b64), c => c.charCodeAt(0)));
  await w.close();
}"""

ANALYZE = """([values, opts]) => {
  const r = SPC.analyze({name: 't', values, labels: values.map((_, i) => 'p' + i), cols: {}}, opts || {});
  return {cl: r.limits.CL.values[0], ucl: r.limits.UCL.values[0], lcl: r.limits.LCL.values[0], flags: r.flags,
          events: r.events.map(e => [e.rule, e.start, e.end]), range: r.range, stats: r.stats, points: r.points.length};
}"""


def workbook(rows, sheet='Data'):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def respell(data, rename):
    """The same .xlsx with its zip entries renamed by rename(name) -> name and text rewritten (other exporters)."""
    src, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            body = src.read(info.filename)
            new, body = rename(info.filename, body)
            zi = zipfile.ZipInfo(new, info.date_time)
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, body)
    return out.getvalue()


def ods(rows):
    """A small OpenDocument spreadsheet (LibreOffice's own .ods layout)."""
    def cell(v):
        if v is None:
            return '<table:table-cell/>'
        if isinstance(v, datetime.datetime):
            return '<table:table-cell office:value-type="date" office:date-value="%s"><text:p>x</text:p></table:table-cell>' % v.isoformat()
        if isinstance(v, (int, float)):
            return '<table:table-cell office:value-type="float" office:value="%r"><text:p>%r</text:p></table:table-cell>' % (v, v)
        return '<table:table-cell office:value-type="string"><text:p>%s</text:p></table:table-cell>' % v
    body = ''.join('<table:table-row>' + ''.join(cell(v) for v in r) + '<table:table-cell table:number-columns-repeated="1000"/></table:table-row>' for r in rows)
    content = ('<?xml version="1.0" encoding="UTF-8"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
               'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
               '<office:body><office:spreadsheet><table:table table:name="Data">' + body +
               '<table:table-row table:number-rows-repeated="1048000"><table:table-cell table:number-columns-repeated="1024"/></table:table-row>'
               '</table:table></office:spreadsheet></office:body></office:document-content>')
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w') as z:
        z.writestr('mimetype', 'application/vnd.oasis.opendocument.spreadsheet')
        z.writestr('content.xml', content, zipfile.ZIP_DEFLATED)
    return out.getvalue()


def long_rows(n, extra=()):
    """Two items in one value column (항목 · 측정값), spec columns on every row."""
    t0 = datetime.datetime(2026, 10, 1, 8, 0)
    rows = [['날짜', '항목', '측정값', 'USL', 'LSL']]
    for i in range(n):
        d = t0 + datetime.timedelta(hours=i)
        rows.append([d, '두께', 10 + (0.05 if i % 2 else -0.05), 10.5, 9.5])
        rows.append([d, '폭', 25 + (0.1 if i % 3 else -0.1), 26, 24])
    rows.extend(extra)
    return rows


# Header of the SPC system export the owner attached to issue #22 (values below are made up).
SPC_HEADERS = ('회사|라인|장비타입|PLAN_NAME|PLAN_DESC|Item|SPEC_NAME|플랜트|공정|PROC_NAME|패키지|DIMENSION|LEAD|STRIP_MARK|고객사|디바이스|'
               'PART_NAME|장비|장비명|LOT_ID|수정자|수정일시|USR_CMF_01|RESULT|USL|TARGET|LSL|EXT_MV|EXT_SIGMA|EXT_MIN|EXT_MAX|EXT_MEDIAN|'
               'EXT_AVERAGE|SITE_QTY|VALUE_COUNT|DISPLAY_FLAG|EXT_MV_UCL|EXT_MV_CENTER|EXT_MV_LCL|EXT_RANGE_UCL|EXT_RANGE_CENTER|EXT_RANGE_LCL|'
               'EXT_RAW_UCL|EXT_RAW_CENTER|EXT_RAW_LCL|TRAN_TIME|UNIT_NUM|DATA_SEQ|SPEC_SEQ|EXT_EWMA_M|EXT_EWMA_S|EXT_EWMA_R|EXT_MA|EXT_MS|'
               'DATA_DESC|DATA_INTERN_DESC|ACCEPT_FLAG|SPEC_LIMIT_ENABLE|RV_STORE_FLAG|LIMIT_ENABLE|IS_NULLVALUE|EXT_SIGMA_LCL|EXT_SIGMA_CENTER|'
               'EXT_SIGMA_UCL|EXT_EWMA_MV_LCL|EXT_EWMA_MV_CENTER|EXT_EWMA_MV_UCL|EXT_MA_MV_LCL|EXT_MA_MV_CENTER|EXT_MA_MV_UCL|EWMA_S_LCL|'
               'EWMA_S_CENTER|EWMA_S_UCL|EWMA_R_LCL|EWMA_R_CENTER|EWMA_R_UCL|EXT_MS_MV_LCL|EXT_MS_MV_CENTER|EXT_MS_MV_UCL|CL_ORIGIN|CPK|'
               'DATA_SEQ1|UNIT_NUM1|' + '|'.join('VALUE_%d' % i for i in range(1, 26)) + '|CNT').split('|')
SPC_SWING = [0.3, 0.1, 0.4, 0.2, 0.5]
SPC_OUT = 10              # data index above EXT_MV_UCL (OCAP)
SPC_RUN = range(20, 28)   # 8 points above EXT_MV_CENTER: the 7th and 8th break the 7 run
SPC_SIGMA_OUT = 5         # EXT_SIGMA above EXT_SIGMA_UCL


def spc_sigma(i, out, run):
    if out and i == SPC_SIGMA_OUT:
        return 2.5
    if run and i in SPC_RUN:
        return 0.7
    if run and i in (SPC_RUN[0] - 1, SPC_RUN[-1] + 1):
        return 0.45
    return 0.45 if i % 2 else 0.68  # alternating around EXT_SIGMA_CENTER 0.564: no SIGMA run


def spc_values():
    vals = [round(6.52 + (-1) ** i * SPC_SWING[i % 5], 2) for i in range(36)]  # alternating around EXT_MV_CENTER 6.52
    vals[SPC_OUT] = 8.9
    vals[SPC_RUN[0] - 1] = vals[SPC_RUN[-1] + 1] = 6.2
    for k, i in enumerate(SPC_RUN):
        vals[i] = 6.9 + 0.03 * (k % 3)
    return vals


def spc_rows(sigma_out=False, sigma_run=False):
    """sigma_out: EXT_SIGMA above EXT_SIGMA_UCL at SPC_SIGMA_OUT; sigma_run: EXT_SIGMA above its center on the SPC_RUN rows too."""
    t0 = datetime.datetime(2026, 7, 1, 2, 6, 2)
    rows = [SPC_HEADERS]
    for i, v in enumerate(spc_values()):
        r = dict.fromkeys(SPC_HEADERS)
        r.update({'회사': 'K5', '라인': 'AS11100', 'PLAN_NAME': 'PLAN-A', 'Item': 'I02214', 'SPEC_NAME': 'TSV-COPLANARITY', '공정': 6311,
                  'DIMENSION': 12, 'LEAD': 134, '장비명': 'AOI-16', 'LOT_ID': 'LOT%03d.C0W' % i, '수정일시': t0 + datetime.timedelta(hours=i),
                  'USL': 13, 'LSL': 0, 'EXT_MV': v, 'EXT_SIGMA': spc_sigma(i, sigma_out, sigma_run),
                  'EXT_SIGMA_LCL': 0.0, 'EXT_SIGMA_CENTER': 0.564, 'EXT_SIGMA_UCL': 2.279, 'EXT_AVERAGE': v, 'SITE_QTY': 25, 'VALUE_COUNT': 25,
                  'DISPLAY_FLAG': 'Y', 'EXT_MV_UCL': 8.234, 'EXT_MV_CENTER': 6.52, 'EXT_MV_LCL': 4.805, 'TRAN_TIME': t0 + datetime.timedelta(hours=i),
                  'UNIT_NUM': 1, 'DATA_SEQ': 8153366 + i, 'SPEC_SEQ': 129588, 'ACCEPT_FLAG': 'YYYYYNYYYY', 'CPK': 3.2, 'VALUE_1': v, 'CNT': 514})
        rows.append([r[h] for h in SPC_HEADERS])
    return rows


SAVE_PICKER = """window.showSaveFilePicker = async (o) => {
  window.__savedAs = o.suggestedName;
  const root = await navigator.storage.getDirectory();
  return await root.getFileHandle('saved_copy.xlsx', {create: true});
};"""

GET_FILE = """async (name) => {
  let root = await navigator.storage.getDirectory();
  const parts = name.split('/');  // 'folder/file' for a file in a sub folder (#32)
  while (parts.length > 1) root = await root.getDirectoryHandle(parts.shift());
  const f = await (await root.getFileHandle(parts[0])).getFile();
  const u8 = new Uint8Array(await f.arrayBuffer());
  let s = ''; for (let i = 0; i < u8.length; i += 8192) s += String.fromCharCode.apply(null, u8.subarray(i, i + 8192));
  return btoa(s);
}"""


class QuietHandler(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


class BuiltFileTest(unittest.TestCase):
    def test_built_file_is_current_and_self_contained(self):
        result = subprocess.run([sys.executable, str(REPO / 'tools' / 'build_html.py'), '--check', '--program', 'SPC_Chart_HTML'], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        html = (PROGRAM / PAGE).read_text(encoding='utf-8')
        self.assertIn("connect-src 'none'", html)
        self.assertNotIn(' src="', html.split('<body>', 1)[1])  # nothing loaded from elsewhere
        self.assertNotIn('href="http', html)
        self.assertIn('<title>%s</title>' % VERSION['title'], html)
        self.assertIn('개발자 : 김민규', html)


@unittest.skipIf(sync_playwright is None, 'playwright / openpyxl not installed')
class SpcChartTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.httpd = http.server.ThreadingHTTPServer(('127.0.0.1', 0), functools.partial(QuietHandler, directory=str(PROGRAM)))
        threading.Thread(target=cls.httpd.serve_forever, daemon=True).start()
        cls.url = 'http://127.0.0.1:%d/%s' % (cls.httpd.server_address[1], PAGE)
        cls.pw = sync_playwright().start()
        cls.browser = cls.pw.chromium.launch()

    @classmethod
    def tearDownClass(cls):
        cls.browser.close()
        cls.pw.stop()
        cls.httpd.shutdown()

    def setUp(self):
        # the page forbids eval (CSP); Playwright's waits need it
        self.context = self.browser.new_context(viewport={'width': 1400, 'height': 900}, bypass_csp=True)
        self.page = self.context.new_page()
        self.errors = []
        self.page.on('pageerror', lambda exc: self.errors.append(str(exc)))
        self.page.on('console', lambda m: self.errors.append(m.text) if m.type == 'error' else None)
        self.page.add_init_script(PICKER)
        self.page.add_init_script(SAVE_PICKER)
        self.page.goto(self.url)

    def tearDown(self):
        self.context.close()
        self.assertEqual(self.errors, [])

    def js(self, expr, arg=None):
        return self.page.evaluate(expr, arg) if arg is not None else self.page.evaluate(expr)

    def put(self, name, data):
        self.js(PUT_FILE, [name, base64.b64encode(data).decode()])

    def open_file(self, name, data):
        self.put(name, data)
        self.js('n => { window.__pick = n; }', name)
        self.page.click('#openBtn')
        self.page.wait_for_function('() => window.__spc.results().length > 0')

    def sim(self, button):
        """[시뮬레이션] → one of its three buttons (#31): 행 삭제만 #delBtn, 행 수정만 #fixAllBtn, 최소 행 변경 #mixBtn"""
        if not self.page.locator(button).is_visible():
            self.page.click('#simBtn')
        self.page.click(button)

    def results(self):
        return self.js("""() => window.__spc.results().map(r => ({name: r.name, n: r.points.length, last: r.x[r.x.length - 1],
            usl: r.limits.USL.values && r.limits.USL.values[0], src: r.limits.UCL.source, events: r.events.map(e => e.rule),
            label: r.points[0] && r.points[0].label}))""")

    # engine
    def test_imr_limits(self):
        r = self.js(ANALYZE, [[10, 12, 11, 13, 12]])
        mr = (2 + 1 + 2 + 1) / 4
        self.assertAlmostEqual(r['cl'], 11.6)
        self.assertAlmostEqual(r['ucl'], 11.6 + 3 * mr / 1.128, places=9)
        self.assertAlmostEqual(r['lcl'], 11.6 - 3 * mr / 1.128, places=9)
        self.assertAlmostEqual(r['range']['ucl'], 3.267 * mr)
        self.assertEqual(r['events'], [])

    def test_xbar_r_limits(self):
        vals = [10, 11, 12, 11, 10, 12, 13, 11, 12, 12]
        r = self.js(ANALYZE, [vals, {'subgroup': 5}])
        self.assertEqual(r['points'], 2)
        rbar = (2 + 2) / 2
        self.assertAlmostEqual(r['cl'], sum(vals) / 10)
        self.assertAlmostEqual(r['ucl'], sum(vals) / 10 + 0.577 * rbar)
        self.assertAlmostEqual(r['range']['ucl'], 2.114 * rbar)

    def test_run_of_seven(self):
        base = [10, 9, 11, 10.5, 9.5, 10, 10.2, 9.8] * 3
        six = self.js(ANALYZE, [base + [10.3] * 6, {'baseline': len(base)}])
        self.assertNotIn('R2', six['flags'])
        seven = self.js(ANALYZE, [base + [10.3] * 7, {'baseline': len(base)}])
        self.assertEqual(seven['flags']['R2'], [len(base) + 6])
        five = self.js(ANALYZE, [base + [10.3] * 5, {'baseline': len(base), 'n': {'R2': 5}}])
        self.assertEqual(five['flags']['R2'], [len(base) + 4])

    def test_trend_of_six_is_one_event(self):  # Trend 6 points by default since v11 (#27); off for XBAR since v13 (#29)
        vals = [10, 9, 11, 10, 9.5, 10.5, 10] * 3 + [9.1, 9.3, 9.5, 9.7, 9.9, 10.1, 10.3, 10.4]
        self.assertNotIn('R3', self.js(ANALYZE, [vals])['flags'])
        r = self.js(ANALYZE, [vals, {'rules': {'R3': True}}])
        n = len(vals)
        self.assertEqual(r['flags']['R3'], [n - 3, n - 2, n - 1])  # 6th, 7th and 8th rising point
        self.assertIn(['R3', n - 3, n - 1], r['events'])

    def test_beyond_limits_and_spec(self):
        vals = [10, 10.1, 9.9, 10, 10.1, 9.9, 10, 10.05, 9.95, 10] * 2 + [11]
        r = self.js("""(vals) => { const n = vals.length;
            const res = SPC.analyze({name: 't', values: vals, labels: vals.map(String), cols: {USL: vals.map(() => 10.8), LSL: [9.2]}}, {baseline: n - 1});
            return {flags: res.flags, usl: res.limits.USL.values[n - 1], lsl: res.limits.LSL.values[n - 1], cpk: res.stats.cpk}; }""", vals)
        self.assertEqual(r['flags']['R1'], [len(vals) - 1])
        self.assertEqual(r['flags']['OOS'], [len(vals) - 1])
        self.assertEqual((r['usl'], r['lsl']), (10.8, 9.2))  # one LSL cell carries to every row
        self.assertIsNotNone(r['cpk'])

    def test_two_of_three_beyond_two_sigma(self):
        vals = [10, 10.2, 9.8, 10.1, 9.9, 10, 10.2, 9.8, 10.1, 9.9]
        r = self.js("""(vals) => { const res = SPC.analyze({name: 't', values: vals.concat([10.75, 10, 10.8]), labels: [], cols: {}},
            {manual: {CL: 10, UCL: 11, LCL: 9}, rules: Object.assign({}, SPC.DEFAULT_ON, {R4: true})}); return res.flags; }""", vals)
        self.assertEqual(r.get('R4'), [12])  # off by default since v8 (#25), turned on here
        self.assertNotIn('R1', r)
        off = self.js("""(vals) => SPC.analyze({name: 't', values: vals.concat([10.75, 10, 10.8]), labels: [], cols: {}},
            {manual: {CL: 10, UCL: 11, LCL: 9}}).flags""", vals)
        self.assertNotIn('R4', off)

    # file reading and live refresh
    def test_long_format_xlsx_with_dates_and_spec(self):
        self.open_file('data.xlsx', workbook(long_rows(30)))
        res = {r['name']: r for r in self.results()}
        self.assertEqual(sorted(res), ['두께', '폭'])
        self.assertEqual(res['두께']['n'], 30)
        self.assertEqual(res['두께']['usl'], 10.5)
        self.assertEqual(res['두께']['label'], '2026-10-01 08:00')
        self.assertEqual(res['두께']['events'], [])
        t = self.js('() => { const t = window.__spc.table(); return [t.label, t.item, t.values]; }')
        self.assertEqual(t, ['날짜', '항목', ['측정값']])
        self.assertEqual(self.page.locator('.item-card').count(), 2)
        self.assertTrue(self.js("""() => { const c = document.querySelector('.cx');
            const d = c.getContext('2d').getImageData(0, 0, c.width, c.height).data;
            for (let i = 0; i < d.length; i += 4) if (d[i + 2] > 200 && d[i] < 80) return true; return false; }"""))  # blue data line drawn

    def test_saving_the_workbook_updates_the_chart_and_raises_alarm(self):
        self.open_file('live.xlsx', workbook(long_rows(30)))
        self.assertTrue(self.page.locator('#alarm').is_hidden())
        last = datetime.datetime(2026, 10, 3, 0, 0)
        self.put('live.xlsx', workbook(long_rows(30, [[last, '두께', 10.9, 10.5, 9.5]])))  # "Ctrl+S" in Excel
        self.page.click('#reloadBtn')  # no auto refresh since v13 (#29)
        self.page.wait_for_function("() => window.__spc.results().find(r => r.name === '두께').points.length === 31", timeout=8000)
        res = {r['name']: r for r in self.results()}
        self.assertAlmostEqual(res['두께']['last'], 10.9)
        self.assertIn('OOS', res['두께']['events'])
        self.assertTrue(self.page.locator('#alarm').is_visible())
        self.assertIn('두께', self.page.inner_text('#alarmText'))
        self.assertEqual(self.page.locator('#events tr.new').count(), len(res['두께']['events']))
        # OCAP: mark the event done
        self.page.click('#alarmOcap')
        self.assertIn('Hold', self.page.inner_text('.steps'))
        self.page.check('#dDone')
        self.page.fill('#dNote', '재측정 OK')
        self.page.click('#dOk')
        self.assertIn('재측정 OK', self.page.inner_text('#events'))

    def test_broken_save_keeps_last_chart(self):
        self.open_file('b.xlsx', workbook(long_rows(10)))
        self.put('b.xlsx', b'PK\x03\x04 half written')
        self.page.click('#reloadBtn')
        self.page.wait_for_function("() => document.getElementById('status').className.includes('err')", timeout=8000)
        self.assertEqual(len(self.results()), 2)
        self.put('b.xlsx', workbook(long_rows(12)))
        self.page.click('#reloadBtn')
        self.page.wait_for_function("() => window.__spc.results()[0].points.length === 12", timeout=8000)

    def test_wide_format_with_limit_columns(self):
        rows = [['Lot', 'A', 'B', 'UCL', 'LCL']] + [['L%02d' % i, 1 + 0.01 * (i % 4), 2 + 0.02 * (i % 3), 1.2, 0.8] for i in range(20)]
        self.open_file('wide.xlsx', workbook(rows))
        res = {r['name']: r for r in self.results()}
        self.assertEqual(sorted(res), ['A', 'B'])
        self.assertEqual(res['A']['src'], '엑셀 열')
        self.assertEqual(res['A']['label'], 'L00')

    def test_csv_in_korean_excel_encoding(self):
        text = '일자,두께\n' + ''.join('2026-10-%02d,%s\n' % (i + 1, 10 + 0.01 * (i % 5)) for i in range(15))
        self.open_file('k.csv', text.encode('cp949'))
        res = self.results()
        self.assertEqual([r['name'] for r in res], ['두께'])
        self.assertEqual(res[0]['n'], 15)

    def test_manual_limits(self):
        self.open_file('m.xlsx', workbook(long_rows(20)))
        self.page.locator('.item-card').first.locator('.lim-btn').click()
        self.page.fill('[data-k="UCL"]', '10.02')
        self.page.click('#dOk')
        r = self.js("() => { const r = window.__spc.results()[0]; return [r.limits.UCL.values[0], r.limits.UCL.source, r.events.length]; }")
        self.assertEqual(r[:2], [10.02, '직접 입력'])
        self.assertGreater(r[2], 0)

    def test_demo(self):
        self.page.click('#demoBtn')
        self.page.wait_for_function('() => window.__spc.results().length === 2')
        self.assertEqual(self.page.locator('.item-card').count(), 2)
        self.assertGreater(self.page.locator('#events tr[data-key]').count(), 0)

    # issue #22: SPC system format, data table, double-click, recommended values, saved copy
    def test_spc_system_format_is_recognised(self):
        self.open_file('20261019214451.xlsx', workbook(spc_rows()))
        t = self.js('() => { const t = window.__spc.table(); return [t.spcForm, t.values, t.label, t.item, t.roles]; }')
        self.assertEqual(t[:4], [True, ['EXT_MV'], 'TRAN_TIME', 'Item'])
        self.assertEqual(t[4], {'USL': 'USL', 'LSL': 'LSL', 'UCL': 'EXT_MV_UCL', 'LCL': 'EXT_MV_LCL', 'CL': 'EXT_MV_CENTER'})
        res = self.results()
        self.assertEqual([r['name'] for r in res], ['I02214'])
        self.assertEqual(res[0]['src'], '엑셀 열')
        self.assertEqual(res[0]['label'], '2026-07-01 02:06')
        yellow = self.js('() => window.__spc.yellow()')
        run_rows = [str(i + 2) for i in SPC_RUN][6:]
        self.assertEqual(sorted(yellow, key=int), [str(SPC_OUT + 2)] + run_rows)  # Excel row = data index + 2
        self.assertEqual(yellow[str(SPC_OUT + 2)], ['R1'])
        self.assertEqual(yellow[run_rows[0]], ['R2'])
        self.assertEqual(self.page.locator('#dataTable tbody tr.y').count(), 3)
        self.assertEqual(self.page.locator('#dataTable tbody tr').count(), 36)
        bg = self.js("() => getComputedStyle(document.querySelector('#dataTable tr.y td')).backgroundColor")
        self.assertEqual(bg, 'rgb(255, 245, 157)')
        self.page.check('#flagOnly')
        self.assertEqual(self.page.locator('#dataTable tbody tr').count(), 3)

    # issue #23: XBAR + SIGMA charts with UCL/LCL lines only, histogram, date and time on the X axis
    def test_spc_system_format_draws_xbar_and_sigma(self):
        self.open_file('x.xlsx', workbook(spc_rows(sigma_out=True)))
        r = self.js("""() => { const r = window.__spc.results()[0];
            return {sigma: r.sigma.values.slice(0, 7), ucl: r.sigma.ucl[0], lcl: r.sigma.lcl[0], rflags: r.rangeFlags, lots: r.lots.slice(0, 1),
                    card: document.querySelector('.item-card').classList.contains('spc'), stats: document.querySelector('.istats').textContent}; }""")
        self.assertEqual(r['sigma'][SPC_SIGMA_OUT], 2.5)
        self.assertEqual(r['sigma'][0], 0.68)
        self.assertEqual((r['ucl'], r['lcl']), (2.279, 0.0))
        self.assertEqual(r['rflags'], [SPC_SIGMA_OUT])
        self.assertEqual(r['lots'], ['LOT000.C0W'])
        self.assertTrue(r['card'])
        self.assertIn('XBAR · SIGMA', r['stats'])
        # the XBAR chart draws only UCL/LCL: USL 13 / LSL 0 do not stretch its scale
        dom = self.js("""() => { const h = document.querySelector('.cx')._hit; return [h.Y(8.234), h.Y(4.805), h.Y(13)]; }""")
        self.assertGreater(dom[1] - dom[0], 150)
        self.assertLess(dom[2], 0)
        # hover the SIGMA point: tooltip with LOT, SIGMA, limits
        self.page.locator('.item-card .cr').first.scroll_into_view_if_needed()
        xy = self.js("""(i) => { const c = document.querySelector('.cr'), h = c._hit, b = c.getBoundingClientRect();
            return [b.left + h.X(i), b.top + h.Y(h.res.sigma.values[i])]; }""", SPC_SIGMA_OUT)
        self.page.mouse.move(xy[0], xy[1])
        tip = self.page.locator('#tip').text_content()
        self.assertIn('SIGMA 2.5', tip)
        self.assertIn('LOT LOT005.C0W', tip)
        self.assertIn('2026-07-01 07:06', tip)
        # double-click on the SIGMA chart selects the row too
        self.page.mouse.dblclick(xy[0], xy[1])
        self.assertEqual(self.js('() => window.__spc.selected()'), SPC_SIGMA_OUT + 2)
        # histogram: frequencies add up to the points, spec limits and ±3σ lines
        h = self.js("() => { const c = document.querySelector('.ch'); return [c._bins, c._marks]; }")
        self.assertEqual(sum(h[0]), 36)
        self.assertEqual(h[1], ['LSL', 'USL', '-3σ', '+3σ'])

    def test_date_time_labels_on_two_lines(self):
        self.open_file('x.xlsx', workbook(spc_rows()))
        texts = self.js("""() => { const out = [], c = document.querySelector('.cx'), ctx = c.getContext('2d'), f = ctx.fillText;
            ctx.fillText = function (t) { out.push(String(t)); return f.apply(this, arguments); };
            window.__spc.refresh(true); return new Promise(r => setTimeout(() => { ctx.fillText = f; r(out); }, 300)); }""")
        self.assertIn('2026-07-02', texts)
        self.assertIn('11:06', texts)  # labels counted from the first point since #24 (02:06 + 3 h steps)
        self.assertIn('XBAR 차트', texts)
        self.assertNotIn('CL', [t.split(' ')[0] for t in texts])
        self.assertNotIn('USL', [t.split(' ')[0] for t in texts])

    def test_double_click_point_goes_to_its_row(self):
        self.open_file('d.xlsx', workbook(spc_rows()))
        self.page.locator('.item-card .cx').first.scroll_into_view_if_needed()
        xy = self.js("""(i) => { const c = document.querySelector('.cx'), h = c._hit, r = c.getBoundingClientRect();
            return [r.left + h.X(i), r.top + h.Y(h.res.x[i])]; }""", SPC_OUT)
        self.page.mouse.dblclick(xy[0], xy[1])
        self.assertEqual(self.js('() => window.__spc.selected()'), SPC_OUT + 2)
        tr = self.page.locator('#dataTable tr.sel')
        self.assertEqual(tr.get_attribute('data-row'), str(SPC_OUT + 2))
        self.assertTrue(self.js("""() => { const r = document.querySelector('#dataTable tr.sel').getBoundingClientRect();
            return r.top >= 0 && r.bottom <= window.innerHeight; }"""))
        self.assertFalse(self.page.locator('#fixOneBtn').is_disabled())

    def test_recommended_values_and_saved_copy(self):
        data = workbook(spc_rows())
        self.open_file('src.xlsx', data)
        mv = SPC_HEADERS.index('EXT_MV')
        # one row: the point above UCL
        self.page.click('#dataTable tr[data-row="%d"] td.rn' % (SPC_OUT + 2))
        self.page.click('#fixOneBtn')
        edits = self.js('() => Object.values(window.__spc.edits())')
        self.assertEqual(len(edits), 1)
        e = edits[0]
        self.assertEqual((e['row'], e['col'], e['orig']), (SPC_OUT + 2, mv, 8.9))
        self.assertLessEqual(abs(e['value'] - 6.52), 2 * (8.234 - 6.52) / 3 + 1e-9)  # inside ±2σ
        self.assertNotIn(str(SPC_OUT + 2), self.js('() => window.__spc.nowFlags()'))
        self.assertEqual(self.page.locator('#dataTable td.fixed').count(), 2)  # the value column after 판정 and the same cell in place
        self.assertEqual(self.page.locator('#dataTable tr.y').count(), 3)  # the row stays yellow
        # all: one value clears the 7 run
        self.sim('#fixAllBtn')
        self.assertEqual(self.js('() => window.__spc.nowFlags()'), {})
        edits = self.js('() => Object.values(window.__spc.edits())')
        self.assertEqual(len(edits), 2)
        flags = self.js('() => window.__spc.results()[0].flags')
        for rule in ('OOS', 'R1', 'R2', 'R3'):
            self.assertNotIn(rule, flags)
        self.assertEqual(self.js('() => window.__spc.simMode()'), 'fix')  # [시뮬레이션 · 행 수정만] (#31), from the file's own values
        self.assertTrue(self.page.locator('#fixOneBtn').is_disabled())
        self.assertIn('시뮬레이션 · 행 수정만', self.page.locator('#delBox').text_content())
        # save the copy
        self.page.click('#saveBtn')
        self.page.wait_for_function("() => document.getElementById('status').textContent.includes('수정본 저장')")
        self.assertRegex(self.js('() => window.__savedAs'), r'^src_SPC수정_\d{8}_\d{4}\.xlsx$')
        saved = base64.b64decode(self.js(GET_FILE, 'saved_copy.xlsx'))
        self.assertEqual(base64.b64decode(self.js(GET_FILE, 'src.xlsx')), data)  # the original is untouched
        ws = openpyxl.load_workbook(io.BytesIO(saved)).active
        orig = openpyxl.load_workbook(io.BytesIO(data)).active
        changed = {(x['row'], x['col']): x['value'] for x in edits}
        for row in range(1, ws.max_row + 1):
            for col in range(1, len(SPC_HEADERS) + 1):
                c = ws.cell(row, col)
                if (row, col - 1) in changed:
                    self.assertEqual(c.value, changed[(row, col - 1)])
                    self.assertEqual(c.fill.fgColor.rgb, 'FFFF0000')
                else:
                    self.assertEqual(c.value, orig.cell(row, col).value)
        for row in map(int, self.js('() => window.__spc.yellow()')):
            self.assertEqual(ws.cell(row, 1).fill.fgColor.rgb, 'FFFFFF00')
        self.assertNotEqual(ws.cell(3, 1).fill.fgColor.rgb, 'FFFFFF00')
        # 행 수정만 off: back to the one value from before; then undo
        self.sim('#fixAllBtn')
        self.assertEqual(len(self.js('() => window.__spc.edits()')), 1)
        self.page.click('#undoBtn')
        self.assertEqual(self.js('() => window.__spc.edits()'), {})
        self.assertEqual(len(self.js('() => window.__spc.nowFlags()')), 3)

    def test_edits_follow_live_refresh(self):
        rows = spc_rows()
        self.open_file('live2.xlsx', workbook(rows))
        self.sim('#fixAllBtn')
        self.assertEqual(len(self.js('() => window.__spc.edits()')), 2)
        rows[SPC_OUT + 1][SPC_HEADERS.index('EXT_MV')] = 6.5  # fixed in Excel and saved
        self.put('live2.xlsx', workbook(rows))
        self.page.click('#reloadBtn')
        self.page.wait_for_function('() => Object.keys(window.__spc.edits()).length === 1', timeout=8000)

    def test_html_table_saved_as_xls(self):
        rows = spc_rows()
        html = '<html><body><table>' + ''.join('<tr>' + ''.join('<td>%s</td>' % ('' if v is None else v) for v in r) + '</tr>' for r in rows) + '</table></body></html>'
        self.open_file('export.xls', html.encode('utf-8'))
        self.assertEqual(self.js('() => window.__spc.table().spcForm'), True)
        self.assertEqual(len(self.js('() => window.__spc.yellow()')), 3)
        self.sim('#fixAllBtn')
        saved = bytes(self.js('() => Array.from(window.__spc.copyBytes())'))
        ws = openpyxl.load_workbook(io.BytesIO(saved)).active
        self.assertEqual(ws.cell(1, 1).value, '회사')
        self.assertEqual(ws.cell(SPC_OUT + 2, SPC_HEADERS.index('EXT_MV') + 1).fill.fgColor.rgb, 'FFFF0000')
        self.assertEqual(ws.cell(SPC_OUT + 2, 1).fill.fgColor.rgb, 'FFFFFF00')
        self.assertTrue(self.js("() => window.__spc.copyName().endsWith('.xlsx')"))

    def test_fix_all_engine_breaks_a_run_with_one_value(self):
        vals = [10, 9, 11, 10.5, 9.5, 10, 10.2, 9.8] * 3 + [10.3] * 9
        r = self.js("""(vals) => { const it = {name: 't', values: vals, labels: vals.map(String), cols: {}};
            const opts = {baseline: 24}, out = SPC.fixAll(it, opts, SPC.FIX_RULES);
            const v = vals.slice(); out.changes.forEach(c => { v[c.row] = c.value; });
            return {out, flags: SPC.analyze(Object.assign({}, it, {values: v}), opts).flags}; }""", vals)
        self.assertEqual(len(r['out']['changes']), 1)
        self.assertEqual(r['out']['changes'][0]['row'], 30)  # the 7th point of the run
        self.assertNotIn('R2', r['flags'])

    # the 7th point of a run cannot move alone: across CL it starts a 7-point trend with the next point
    def test_run_point_that_only_another_run_point_can_fix(self):
        vals = [6.9, 6.2, 6.8, 6.3, 6.7, 6.25] * 3 + [6.49, 6.74, 7.56, 5.78, 6.58, 6.93, 6.91, 6.69, 6.62, 6.56, 6.91, 5.04, 6.94, 6.59, 6.82, 6.28, 6.85]
        r = self.js("""(vals) => { const n = vals.length, c = v => vals.map(() => v);
            const it = {name: 't', values: vals, labels: vals.map(String), cols: {UCL: c(8.234), CL: c(6.52), LCL: c(4.805)}};
            const opts = {rules: Object.assign({}, SPC.DEFAULT_ON, {R3: true})}, rules = SPC.FIX_RULES, row = 28;  // Trend on (off for XBAR since v13)
            const out = SPC.fixAll(it, opts, rules), v = vals.slice(); out.changes.forEach(ch => { v[ch.row] = ch.value; });
            return {flags: SPC.analyze(it, opts).flags, own: SPC.suggest(it, opts, row, rules), change: SPC.suggestChange(it, opts, row, rules),
                    out, after: SPC.analyze(Object.assign({}, it, {values: v}), opts).flags}; }""", vals)
        self.assertEqual(r['flags']['R2'], [28])
        self.assertIsNone(r['own'])  # 6.91 -> any value: still R2 above CL, or R3 (200..206 falling) below it
        self.assertIn(r['change']['row'], range(22, 28))  # another point of the run
        self.assertEqual(r['out']['failed'], 0)
        self.assertEqual(len(r['out']['changes']), 1)
        for rule in ('R2', 'R3'):
            self.assertNotIn(rule, r['after'])

    # v8 (#25): 2/3 > 2σ and 4/5 > 1σ are off by default, also for settings saved by an older version
    def test_extra_rules_off_by_default(self):
        self.assertEqual([self.js('() => window.__spc.settings.rules')[k] for k in ('R1', 'R2', 'R3', 'R4', 'R5')], [True, True, False, False, False])
        self.js("() => localStorage.setItem('spc.settings', JSON.stringify({rules: {R4: true, R5: true}, v4: 1}))")
        self.page.reload()
        self.assertEqual([self.js('() => window.__spc.settings.rules')[k] for k in ('R4', 'R5')], [False, False])
        self.open_file('any.xlsx', workbook(long_rows(10)))  # reading a file saves the settings (with the v8 mark)
        self.js("() => { const s = JSON.parse(localStorage.getItem('spc.settings')); s.rules.R4 = true; localStorage.setItem('spc.settings', JSON.stringify(s)); }")
        self.page.reload()  # turned on again by the owner: stays on
        self.assertTrue(self.js('() => window.__spc.settings.rules.R4'))

    # issue #24: a long period scrolls sideways instead of being cut; auto refresh is off by default
    def test_long_period_scrolls_sideways(self):
        self.assertEqual(self.page.locator('#liveChk').count(), 0)  # auto refresh removed in v13 (#29)
        self.open_file('long.xlsx', workbook(long_rows(400)))
        self.assertEqual(self.js('() => window.__spc.settings.window'), 0)
        self.assertEqual(self.page.input_value('#pitchSel'), '16')
        geo = self.js("""() => { const s = document.querySelector('.xscroll'), c = s.querySelector('.cx');
            return {sw: s.scrollWidth, cw: s.clientWidth, canvas: c.clientWidth, from: c._hit.from, to: c._hit.to, n: c._hit.res.points.length}; }""")
        self.assertEqual(geo['n'], 400)  # nothing cut off
        self.assertGreaterEqual(geo['sw'], 400 * 16)
        self.assertEqual(geo['canvas'], geo['cw'])  # the canvas stays the card width
        self.assertEqual(geo['from'], 0)
        self.assertLess(geo['to'], 120)
        # scroll to the end: the last points are drawn, the MR chart follows, the Y axis stays at the left
        self.js("() => { const s = document.querySelector('.xscroll'); s.scrollLeft = s.scrollWidth; }")
        self.page.wait_for_function("() => document.querySelector('.cx')._hit.to === 400")
        h = self.js("""() => { const c = document.querySelector('.cx')._hit, r = document.querySelector('.cr')._hit;
            return [c.from, r.to, c.X(399) < document.querySelector('.cx').clientWidth - 92]; }""")
        self.assertGreater(h[0], 250)
        self.assertEqual(h[1], 400)
        self.assertTrue(h[2])
        # hover / double-click still find the right point after scrolling
        self.page.locator('.item-card .cx').first.scroll_into_view_if_needed()
        xy = self.js("""(i) => { const c = document.querySelector('.cx'), h = c._hit, b = c.getBoundingClientRect();
            return [b.left + h.X(i), b.top + h.Y(h.res.x[i])]; }""", 395)
        self.page.mouse.dblclick(xy[0], xy[1])
        row = self.js('() => window.__spc.selected()')
        self.assertEqual(row, self.js("() => { const r = document.querySelector('.cx')._hit.res; return r.excelRows[r.points[395].rows[0]]; }"))
        # picking a row in the table scrolls the chart back to it
        first = self.js("() => { const r = document.querySelector('.cx')._hit.res; return r.excelRows[r.points[3].rows[0]]; }")
        self.page.locator('#dataTable tr[data-row="%d"]' % first).click()
        self.page.wait_for_function("() => document.querySelector('.cx')._hit.from <= 3")
        # 화면에 맞춤: the whole period on one screen, no scroll
        self.page.select_option('#pitchSel', '0')
        geo = self.js("""() => { const s = document.querySelector('.xscroll'), h = s.querySelector('.cx')._hit; return [s.scrollWidth - s.clientWidth, h.from, h.to]; }""")
        self.assertEqual(geo, [0, 0, 400])

    # issue #25 follow-up: the SIGMA chart is judged too (OCAP, 7 run) and its rows get recommended values with the XBAR ones
    def test_sigma_ocap_and_run_with_joint_recommended_values(self):
        rows = spc_rows(sigma_out=True, sigma_run=True)
        self.open_file('sig.xlsx', workbook(rows))
        mv, sg = SPC_HEADERS.index('EXT_MV'), SPC_HEADERS.index('EXT_SIGMA')
        yellow = self.js('() => window.__spc.yellow()')
        run_rows = [str(i + 2) for i in SPC_RUN][6:]
        self.assertEqual(yellow[str(SPC_SIGMA_OUT + 2)], ['SR1'])
        self.assertEqual(sorted(yellow[run_rows[0]]), ['R2', 'SR2'])  # both charts break the 7 run on this row
        rules = self.js('() => window.__spc.results()[0].events.map(e => e.rule)')
        self.assertIn('SR1', rules)
        self.assertIn('SR2', rules)
        verdict = self.page.locator('#dataTable tr[data-row="%d"] td.verdict' % (SPC_SIGMA_OUT + 2)).text_content()
        self.assertIn('SIGMA 관리한계 이탈', verdict)
        # one row: EXT_MV and EXT_SIGMA together
        self.page.click('#dataTable tr[data-row="%s"] td.rn' % run_rows[0])
        self.page.click('#fixOneBtn')
        edits = {(e['row'], e['col']): e['value'] for e in self.js('() => Object.values(window.__spc.edits())')}
        self.assertEqual(sorted(edits), [(int(run_rows[0]), mv), (int(run_rows[0]), sg)])
        self.assertLess(edits[(int(run_rows[0]), sg)], 0.564)  # across the SIGMA center line
        self.assertGreaterEqual(edits[(int(run_rows[0]), sg)], 0)
        self.assertNotIn(run_rows[0], self.js('() => window.__spc.nowFlags()'))
        # all: no XBAR or SIGMA rule left, no sigma below 0
        self.sim('#fixAllBtn')
        self.assertEqual(self.js('() => window.__spc.nowFlags()'), {})
        self.assertEqual(self.js('() => window.__spc.results()[0].sigmaFlags'), {})
        edits = self.js('() => Object.values(window.__spc.edits())')
        sig = [e for e in edits if e['col'] == sg]
        self.assertEqual(sorted(e['row'] for e in sig), [SPC_SIGMA_OUT + 2, int(run_rows[0])])
        for e in sig:
            self.assertTrue(0 <= e['value'] <= 2.279)
        saved = bytes(self.js('() => Array.from(window.__spc.copyBytes())'))
        ws = openpyxl.load_workbook(io.BytesIO(saved)).active
        c = ws.cell(SPC_SIGMA_OUT + 2, sg + 1)
        self.assertEqual(c.fill.fgColor.rgb, 'FFFF0000')
        self.assertEqual(c.value, [e['value'] for e in sig if e['row'] == SPC_SIGMA_OUT + 2][0])
        # the SIGMA value column is shown next to EXT_MV after 판정
        self.assertEqual(self.page.locator('#dataTable thead th.vcol').first.text_content(), 'EXT_MV')
        self.assertEqual(self.page.locator('#dataTable thead th.vcol').nth(1).text_content(), 'EXT_SIGMA')

    # issue #27: deleting the yellow rows can join two short runs into a new 7 run; the plan goes round by round
    def test_delete_plan_rounds(self):
        vals = [-1, 1, 2, 1, 2, 1, 2, -5, 1, 2, 1, 2, 1, 2, -1, 1, -1]
        r = self.js("""(vals) => { const c = v => vals.map(() => v);
            const it = {name: 't', values: vals, labels: vals.map(String), cols: {UCL: c(3), CL: c(0), LCL: c(-3)}};
            const rules = SPC.FIX_RULES, keys = vals.map((_, k) => k + 2);
            return SPC.deletePlan([{item: it, opts: {}, rules, keys}]); }""", vals)
        self.assertTrue(r['clean'])
        self.assertEqual([[x['key'] for x in rd] for rd in r['rounds']], [[9], [10, 11, 12, 13, 14, 15]])  # -5 (row 9), then points 7..12 of the joined 12-point run
        self.assertEqual(r['rounds'][0][0]['rules'], ['R1'])
        self.assertEqual(r['rounds'][1][0]['rules'], ['R2'])

    # #27 follow-up: one row at a time; an XBAR row's deletion shortens the SIGMA run it sat in, so that run's
    # 7th row stays (deleting a whole round at once took both)
    def test_delete_plan_one_row_at_a_time(self):
        mv = [(-1) ** i for i in range(20)]
        mv[9] = 5
        sg = [0.4 if i % 2 else 0.6 for i in range(20)]
        for i in range(6, 13):
            sg[i] = 0.7
        sg[13] = 0.4
        r = self.js("""([mv, sg]) => { const c = (a, v) => a.map(() => v), keys = mv.map((_, k) => k + 2);
            const x = {name: 'x', values: mv, labels: keys.map(String), cols: {UCL: c(mv, 3), CL: c(mv, 0), LCL: c(mv, -3)}};
            const s = {name: 's', values: sg, labels: keys.map(String), cols: {UCL: c(sg, 2), CL: c(sg, 0.5), LCL: c(sg, 0)}};
            const before = SPC.analyze(s, {rules: {R1: true, R2: true, R3: true}}).flags;
            return {before, plan: SPC.deletePlan([{item: x, opts: {}, rules: SPC.FIX_RULES, keys, tag: ''},
              {item: s, opts: {rules: {R1: true, R2: true, R3: true}, min: 0}, rules: ['R1', 'R2', 'R3'], keys, tag: 'S'}])}; }""", [mv, sg])
        self.assertEqual(r['before']['R2'], [12])  # the SIGMA run's 7th point (row 14) is yellow too
        self.assertTrue(r['plan']['clean'])
        self.assertEqual([[x['key'] for x in rd] for rd in r['plan']['rounds']], [[11]])  # only the XBAR OCAP row

    # v11 (#27): Trend 6 by default; a 7 saved by an older version goes, a count picked on purpose stays
    def test_trend_default_six_and_old_settings(self):
        for saved, want in ((7, 6), (9, 9)):
            self.js("""(n) => localStorage.setItem('spc.settings', JSON.stringify({v4: 1, v8: 1, n: {R2: 7, R3: n}}))""", saved)
            self.page.reload()
            self.open_file('t.xlsx', workbook(long_rows(10)))
            n = self.js('() => window.__spc.settings.n')
            self.assertEqual(n.get('R3', 6), want)

    # v13 (#29): XBAR = OCAP, 7 run; SIGMA = OCAP, 7 run, Trend 6 by default, set apart in [규칙 · OCAP 설정]
    def test_rules_per_chart(self):
        rows = spc_rows()
        mv, sg = SPC_HEADERS.index('EXT_MV'), SPC_HEADERS.index('EXT_SIGMA')
        for k, i in enumerate(range(30, 36)):  # 6 rising points on both charts, no run of 7, inside the limits
            rows[i + 1][mv] = round(6.0 + 0.2 * k, 2)
            rows[i + 1][sg] = round(0.3 + 0.1 * k, 2)
        self.open_file('trend.xlsx', workbook(rows))
        flags = lambda: self.js('() => { const r = window.__spc.results()[0]; return [Object.keys(r.flags), Object.keys(r.sigmaFlags)]; }')
        x, s = flags()
        self.assertNotIn('R3', x)
        self.assertIn('SR3', s)
        self.assertIn('37', self.js('() => window.__spc.yellow()'))  # excel row of the 6th rising SIGMA point
        self.page.click('#settingsBtn')
        self.assertIn('SIGMA 차트', self.page.inner_text('#modal'))
        self.assertEqual(self.page.locator('#dPoll').count(), 0)
        self.page.check('[data-chart="x"][data-rule="R3"]')
        self.page.uncheck('[data-chart="s"][data-rule="R3"]')
        self.page.fill('[data-chart="s"][data-n="R2"]', '5')
        self.page.click('#dOk')
        x, s = flags()
        self.assertIn('R3', x)
        self.assertNotIn('SR3', s)
        st = self.js('() => window.__spc.settings')
        self.assertEqual((st['n'], st['sigma']['n']), ({}, {'R2': 5}))
        self.page.click('#settingsBtn')
        self.page.click('#dReset')
        self.page.click('#dOk')
        x, s = flags()
        self.assertEqual(('R3' in x, 'SR3' in s), (False, True))
        # settings of v12 (one shared set, Trend on, auto refresh) -> XBAR Trend off once, SIGMA defaults, no auto refresh
        self.js("() => localStorage.setItem('spc.settings', JSON.stringify({v4: 1, v8: 1, v11: 1, rules: {R3: true}, auto: true, pollSec: 2}))")
        self.page.reload()
        st = self.js('() => window.__spc.settings')
        self.assertFalse(st['rules']['R3'])
        self.assertEqual(st['sigma']['rules'], {'R1': True, 'R2': True, 'R3': True})
        self.assertNotIn('auto', st)
        self.assertNotIn('pollSec', st)

    def test_delete_simulation_clears_xbar_and_sigma(self):
        rows = spc_rows(sigma_out=True, sigma_run=True)
        self.open_file('del.xlsx', workbook(rows))
        yellow = self.js('() => window.__spc.yellow()')
        self.assertTrue(yellow)
        self.sim('#delBtn')
        plan = self.js('() => window.__spc.delPlan()')
        self.assertTrue(plan['clean'])
        self.assertEqual(sorted(str(x['key']) for x in plan['rounds'][0]), sorted(yellow))  # round 1 = the yellow rows
        self.assertEqual(self.js('() => window.__spc.nowFlags()'), {})
        self.assertEqual(self.js('() => window.__spc.results()[0].sigmaFlags'), {})
        self.assertEqual(self.js('() => window.__spc.results()[0].points.length'), len(rows) - 1 - plan['count'])
        self.assertEqual(self.js('() => window.__spc.yellow()'), yellow)  # still the file's own yellow rows
        first = plan['rounds'][0][0]['key']
        self.assertIn('del', self.page.locator('#dataTable tr[data-row="%d"]' % first).get_attribute('class'))
        self.assertIn('삭제 1차', self.page.locator('#dataTable tr[data-row="%d"] td.verdict' % first).text_content())
        self.assertIn('OCAP · 7run이 나오지 않습니다', self.page.locator('#delBox').text_content())
        self.assertTrue(self.page.locator('#fixOneBtn').is_disabled())
        self.assertIn('✓', self.page.locator('#delBtn').text_content())
        # every round is checked again from the file: no rule left once all of them are gone
        with self.page.expect_download() as dl:
            self.page.click('#delCsvBtn')
        with open(dl.value.path(), encoding='utf-8-sig') as f:
            csv = f.read().splitlines()
        self.assertEqual(csv[0], '차수,엑셀 행,X축,LOT_ID,판정')
        self.assertEqual(len(csv) - 1, plan['count'])
        self.sim('#delBtn')
        self.assertIsNone(self.js('() => window.__spc.delPlan()'))
        self.assertEqual(self.js('() => window.__spc.nowFlags()'), yellow)
        self.assertEqual(self.js('() => window.__spc.results()[0].points.length'), len(rows) - 1)

    # issue #28: the delete rows saved in the OCAP 삭제 양식 (EDC Plan · Spec seq# · Data seq#), one sheet per spec
    def test_ocap_delete_form(self):
        rows = spc_rows(sigma_out=True, sigma_run=True)
        hp, hs, hd = (SPC_HEADERS.index(h) for h in ('PLAN_NAME', 'SPEC_SEQ', 'DATA_SEQ'))
        for k, r in enumerate(rows[1:]):
            r[hp] = 'ITEM-A(TARGET 13 UM)_0/13 : F01(LINE-1)'
            r[hs] = 129588 if k < 15 else 130121  # two specs in one file
        self.open_file('ocap.xlsx', workbook(rows))
        self.sim('#delBtn')
        plan = self.js('() => window.__spc.delPlan()')
        keys = sorted(x['key'] for rd in plan['rounds'] for x in rd)
        self.assertTrue(keys)
        self.page.click('#ocapFormBtn')
        self.page.wait_for_function("() => /OCAP 삭제 양식 저장: /.test(document.body.textContent) && window.__savedAs")
        self.assertRegex(self.js('() => window.__savedAs'), r'^OCAP 삭제 양식_(129588|130121|129588\+130121)_\d{12}\.xlsx$')  # #32: 제목_SPEC_SEQ_yyyymmddhhmm
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.js(GET_FILE, 'saved_copy.xlsx'))))
        self.assertEqual(wb.sheetnames, ['Sheet1'])  # v14 (#30): one sheet like the owner's form file
        ws = wb['Sheet1']
        self.assertEqual(ws['B2'].value, 'OCAP 삭제 양식')
        self.assertTrue(ws['B2'].font.b)
        self.assertIsNone(ws['E1'].value)  # no [Option] box any more
        self.assertEqual([ws.cell(5, c).value for c in range(3, 6)], ['OCAP 전체 삭제 시', 'OCAP 일부 삭제 시', '요청 사항'])
        self.assertEqual([ws.cell(6, c).value for c in range(2, 6)], ['EDC Plan', 'Spec seq#', 'Data seq#', '요청 내용'])
        self.assertEqual(ws['B6'].fill.fgColor.rgb, 'FFAEAAAA')
        self.assertEqual(ws['B6'].border.left.style, 'thin')
        self.assertEqual(round(ws.column_dimensions['B'].width, 1), 64.9)
        for i, k in enumerate(keys):  # one line per row to delete: Data seq# = DATA_SEQ
            r = rows[k - 1]
            self.assertEqual([ws.cell(7 + i, c).value for c in range(2, 5)], [r[hp], r[hs], r[hd]])
            self.assertEqual(ws.cell(7 + i, 5).value, 'Data 자체를 완전 삭제 & 히스토리 삭제 (Data 자체 완전 삭제')
        self.assertIsNone(ws.cell(7 + len(keys), 2).value)
        self.assertEqual(ws.cell(26, 5).border.bottom.style, 'thin')  # empty bordered lines down to row 26
        self.assertEqual(ws.max_row, max(26, 6 + len(keys)))

    # issue #30: the rows with new values in the owner's RAW_DATA form; raw values moved / stretched to the new EXT_MV / EXT_SIGMA
    def test_raw_data_form(self):
        rows = spc_rows(sigma_out=True, sigma_run=True)
        mv, sg, hd = (SPC_HEADERS.index(h) for h in ('EXT_MV', 'EXT_SIGMA', 'DATA_SEQ'))
        v1 = SPC_HEADERS.index('VALUE_1')
        pattern = [(-1) ** j * (1 + j % 4) for j in range(25)]
        pm = sum(pattern) / 25
        psd = (sum((x - pm) ** 2 for x in pattern) / 24) ** 0.5
        for r in rows[1:]:  # 25 raw values whose mean = EXT_MV, sample SD = EXT_SIGMA
            for j in range(25):
                r[v1 + j] = round(r[mv] + (pattern[j] - pm) / psd * r[sg], 6)
        self.open_file('raw.xlsx', workbook(rows))
        self.assertTrue(self.page.locator('#rawFormBtn').is_disabled())
        self.sim('#fixAllBtn')
        edits = self.js('() => Object.values(window.__spc.edits())')
        self.assertTrue(any(e['col'] == mv for e in edits) and any(e['col'] == sg for e in edits))
        self.page.click('#rawFormBtn')
        self.page.wait_for_function("() => /수정 양식 저장: /.test(document.body.textContent) && window.__savedAs")
        self.assertRegex(self.js('() => window.__savedAs'), r'^RAW_DATA 수정 양식_129588_\d{12}\.xlsx$')
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.js(GET_FILE, 'saved_copy.xlsx'))))
        ws = wb['Sheet1']
        self.assertEqual(ws['B2'].value, 'S/S 25>')
        self.assertEqual([c.value for c in ws[3]][1:], ['DATA_SEQ#'] + ['Raw Value %d' % j for j in range(1, 26)] + ['MAX', 'MIN', 'AVG', 'STD'])
        self.assertEqual(ws['B3'].fill.fgColor.rgb, 'FFC6E0B4')
        self.assertEqual(ws['C3'].fill.fgColor.rgb, 'FFD9E1F2')
        self.assertTrue(ws['C3'].font.b)
        self.assertEqual(ws['AD4'].number_format, '###0.0000;\\-###0.0000')
        changed = sorted({e['row'] for e in edits})
        self.assertEqual(ws.max_row, 3 + len(changed))
        new = {(e['row'], e['col']): e['value'] for e in edits}
        for i, row in enumerate(changed):
            orig = rows[row - 1]
            line = [c.value for c in ws[4 + i]]
            self.assertEqual(ws.cell(4 + i, 3).fill.fgColor.rgb, 'FFFFFF00')
            self.assertEqual(line[1], orig[hd])
            raw = line[2:27]
            want_mv, want_sg = new.get((row, mv), orig[mv]), new.get((row, sg), orig[sg])
            m = sum(raw) / 25
            sd = (sum((x - m) ** 2 for x in raw) / 24) ** 0.5
            self.assertAlmostEqual(m, want_mv, places=3)
            self.assertAlmostEqual(sd, want_sg, places=3)
            self.assertAlmostEqual(line[29], m, places=6)  # AVG
            self.assertAlmostEqual(line[30], sd, places=6)  # STD
            self.assertEqual((line[27], line[28]), (max(raw), min(raw)))

    # issue #30: deleting and changing values mixed, fewest rows; never more than deleting only or values only
    def test_min_change_plan_engine(self):
        vals = [-1, 1, 2, 1, 2, 1, 2, -5, 1, 2, 1, 2, 1, 2, -1, 1, -1]
        r = self.js("""(vals) => { const c = v => vals.map(() => v);
            const it = {name: 't', values: vals, labels: vals.map(String), cols: {UCL: c(3), CL: c(0), LCL: c(-3)}};
            const keys = vals.map((_, k) => k + 2), series = [{item: it, opts: {}, rules: SPC.FIX_RULES, keys}];
            const p = SPC.minChangePlan(series), v = vals.slice(), gone = {};
            p.deletes.forEach(x => { gone[x.key] = 1; });
            p.changes.forEach(x => { v[x.index] = x.value; });
            const left = SPC.analyze(Object.assign({}, it, {values: v.map((x, k) => gone[keys[k]] ? null : x)}), {}).flags;
            return {p, left: SPC.FIX_RULES.filter(id => (left[id] || []).length)}; }""", vals)
        p = r['p']
        self.assertTrue(p['clean'])
        self.assertEqual(r['left'], [])
        # deleting only needs 7 rows (#27 test above); a new value on row 9 (-5) alone is enough
        self.assertEqual(p['compare']['delete'], 7)
        self.assertEqual(p['rows'], 1)
        self.assertEqual([c['key'] for c in p['changes']], [9])
        self.assertLessEqual(p['rows'], min(n for n in p['compare'].values() if n is not None))

    def test_min_change_button(self):
        rows = spc_rows(sigma_out=True, sigma_run=True)
        self.open_file('mix.xlsx', workbook(rows))
        yellow = self.js('() => window.__spc.yellow()')
        self.sim('#mixBtn')
        self.page.wait_for_function('() => window.__spc.mixPlan()')
        p = self.js('() => window.__spc.mixPlan()')
        self.assertTrue(p['clean'])
        self.assertEqual(self.js('() => window.__spc.nowFlags()'), {})
        self.assertEqual(self.js('() => window.__spc.results()[0].sigmaFlags'), {})
        self.assertEqual(self.js('() => window.__spc.yellow()'), yellow)
        cmp = [n for n in p['compare'].values() if n is not None]
        self.assertLessEqual(p['rows'], min(cmp))
        self.assertEqual(len(self.js('() => window.__spc.edits()')), len(p['changes']))
        box = self.page.locator('#delBox').text_content()
        self.assertIn('최소 행 변경', box)
        self.assertIn('OCAP · 7run이 나오지 않습니다', box)
        self.assertEqual(self.js('() => window.__spc.simMode()'), 'mix')
        self.assertEqual(self.page.locator('#mixBtn').text_content(), '✓ 최소 행 변경 (삭제+수정)')
        with self.page.expect_download() as dl:
            self.page.click('#delCsvBtn')
        with open(dl.value.path(), encoding='utf-8-sig') as f:
            csv = f.read().splitlines()
        self.assertEqual(csv[0], '구분,엑셀 행,X축,LOT_ID,열,원래 값,새 값,판정')
        self.assertEqual(len(csv) - 1, len(p['deletes']) + len(p['changes']))
        self.sim('#mixBtn')
        self.assertIsNone(self.js('() => window.__spc.mixPlan()'))
        self.assertEqual(self.js('() => window.__spc.edits()'), {})
        self.assertEqual(self.js('() => window.__spc.nowFlags()'), yellow)

    # issue #31: 측정값 열 folded by default; [선택 행 추천값 입력] | [시뮬레이션]; undo / save under the table
    def test_layout_value_fold_and_buttons(self):
        self.open_file('lay.xlsx', workbook(spc_rows()))
        self.assertFalse(self.js("() => document.getElementById('valueFold').open"))
        self.assertFalse(self.page.locator('#valueBox input').first.is_visible())
        self.assertIn('EXT_MV', self.page.locator('#valueSum').text_content())
        self.page.click('#valueFold summary')
        self.assertTrue(self.page.locator('#valueBox input').first.is_visible())
        self.assertIn('접기', self.page.locator('#valueSum').text_content())
        order = self.js("""() => Array.from(document.querySelectorAll('#dataBox .dhead > button, #dataBox .dhead > .vsep, #dataBox .dhead > .simgroup')).map(e => e.id || e.className)""")
        self.assertEqual(order, ['fixOneBtn', 'vsep', 'simgroup'])
        self.assertFalse(self.page.locator('#delBtn').is_visible())  # the three modes sit under [시뮬레이션]
        self.page.click('#simBtn')
        self.assertEqual([self.page.locator(b).text_content() for b in ('#delBtn', '#fixAllBtn', '#mixBtn')], ['행 삭제만', '행 수정만', '최소 행 변경 (삭제+수정)'])
        under = self.js("""() => { const w = document.querySelector('#dataBox .dwrap').getBoundingClientRect();
            return ['undoBtn', 'saveBtn'].map(id => document.getElementById(id).getBoundingClientRect().top >= w.bottom); }""")
        self.assertEqual(under, [True, True])

    # issue #31: 행 수정만 → RAW_DATA 수정 양식 from its box; off again = the edits from before
    def test_fix_mode_saves_raw_form(self):
        self.open_file('fixm.xlsx', workbook(spc_rows(sigma_out=True, sigma_run=True)))
        yellow = self.js('() => window.__spc.yellow()')
        self.sim('#fixAllBtn')
        self.assertEqual(self.js('() => window.__spc.nowFlags()'), {})
        n = len(self.js('() => window.__spc.edits()'))
        self.assertGreater(n, 0)
        self.assertIn('OCAP · 7run이 나오지 않습니다', self.page.locator('#delBox').text_content())
        self.page.click('#fixFormBtn')
        self.page.wait_for_function("() => /수정 양식 저장: /.test(document.body.textContent) && window.__savedAs")
        self.assertRegex(self.js('() => window.__savedAs'), r'^RAW_DATA 수정 양식_129588_\d{12}\.xlsx$')
        ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.js(GET_FILE, 'saved_copy.xlsx'))))['Sheet1']
        self.assertEqual(ws['B3'].value, 'DATA_SEQ#')
        self.assertEqual(ws.max_row, 3 + len({e['row'] for e in self.js('() => Object.values(window.__spc.edits())')}))
        with self.page.expect_download() as dl:
            self.page.click('#delCsvBtn')
        with open(dl.value.path(), encoding='utf-8-sig') as f:
            self.assertEqual(len(f.read().splitlines()) - 1, n)
        self.sim('#fixAllBtn')
        self.assertIsNone(self.js('() => window.__spc.simMode()'))
        self.assertEqual(self.js('() => window.__spc.edits()'), {})
        self.assertEqual(self.js('() => window.__spc.nowFlags()'), yellow)

    # issue #31: 최소 행 변경 saves both forms at once, each in its own file (one folder picker)
    def test_mix_mode_saves_both_forms(self):
        self.js("""() => { window.showDirectoryPicker = async () => navigator.storage.getDirectory(); }""")
        rows = spc_rows(sigma_out=True, sigma_run=True)
        rows[-3][SPC_HEADERS.index('USL')] = 4.0  # out of spec, and no new value fits inside the spec and CL ± 2σ: deleted
        self.open_file('both.xlsx', workbook(rows))
        self.sim('#mixBtn')
        self.page.wait_for_function('() => window.__spc.mixPlan()')
        p = self.js('() => window.__spc.mixPlan()')
        self.assertTrue(p['deletes'] and p['changes'], p)
        self.page.click('#bothFormBtn')
        self.page.wait_for_function("() => /양식 2개 저장: /.test(document.getElementById('status').textContent)")
        # #32: both files in a new folder named SPEC_SEQ, each named 제목_SPEC_SEQ_yyyymmddhhmm
        self.assertIn('\\129588 폴더에', self.page.locator('#status').text_content())
        names = self.js("""async () => { const out = []; const d = await (await navigator.storage.getDirectory()).getDirectoryHandle('129588');
          for await (const k of d.keys()) out.push(k); return out; }""")
        ocap = [n for n in names if re.match(r'^OCAP 삭제 양식_129588_\d{12}\.xlsx$', n)]
        raw = [n for n in names if re.match(r'^RAW_DATA 수정 양식_129588_\d{12}\.xlsx$', n)]
        self.assertEqual((len(ocap), len(raw)), (1, 1), names)
        ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.js(GET_FILE, '129588/' + ocap[0]))))['Sheet1']
        self.assertEqual(ws['B6'].value, 'EDC Plan')
        self.assertEqual(len([r for r in range(7, ws.max_row + 1) if ws.cell(r, 4).value is not None]), len(p['deletes']))
        ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.js(GET_FILE, '129588/' + raw[0]))))['Sheet1']
        self.assertEqual(ws.max_row, 3 + len({c['key'] for c in p['changes']}))

    # issue #32: the file on screen at the top, no [최근 파일 다시 열기] button
    def test_current_file_at_top(self):
        self.assertEqual(self.page.locator('#recentBtn').count(), 0)
        self.assertFalse(self.page.locator('#curFile').is_visible())
        self.open_file('lot_week42.xlsx', workbook(spc_rows()))
        self.assertEqual(self.page.locator('#curFile').text_content(), '보고 있는 파일: lot_week42.xlsx')
        self.assertTrue(self.page.locator('#curFile').is_visible())

    # issue #31: drag over the data table, Ctrl+C: an HTML table with the screen's colours (Excel keeps them) + text
    def test_drag_copy_keeps_colours(self):
        self.open_file('copy.xlsx', workbook(spc_rows()))
        self.page.click('#dataTable tr[data-row="%d"] td.rn' % (SPC_OUT + 2))
        self.page.click('#fixOneBtn')
        row = SPC_OUT + 2
        a = self.page.locator('#dataTable tr[data-row="%d"] td.rn' % (row - 1))
        b = self.page.locator('#dataTable tr[data-row="%d"] td' % row).nth(3)
        a.scroll_into_view_if_needed()
        ba, bb = a.bounding_box(), b.bounding_box()
        self.page.mouse.move(ba['x'] + 5, ba['y'] + 5)
        self.page.mouse.down()
        self.page.mouse.move(bb['x'] + 5, bb['y'] + 5, steps=5)
        self.page.mouse.up()
        self.assertEqual(self.page.locator('#dataTable .pick').count(), 2 * 4)
        self.assertFalse(self.page.locator('#copyBtn').is_disabled())
        out = self.js("""() => { const dt = new DataTransfer(); document.dispatchEvent(new ClipboardEvent('copy', {clipboardData: dt, bubbles: true}));
            return [dt.getData('text/html'), dt.getData('text/plain')]; }""")
        html, text = out
        lines = text.split('\r\n')
        self.assertEqual(len(lines), 2)
        self.assertEqual([len(l.split('\t')) for l in lines], [4, 4])
        self.assertEqual(lines[1].split('\t')[0], str(row))
        self.assertEqual(html.count('<tr>'), 2)
        self.assertIn('background:#fff59d', html)  # the yellow row
        self.assertIn('background:#ef4444', html)  # the red cell (EXT_MV after 판정)
        self.assertIn('복사: 2행 × 4열', self.page.locator('#status').text_content())

    # issue #30: [설정 저장] keeps the settings in the browser and a settings file; [설정 불러오기] reads it back
    def test_save_and_load_settings(self):
        self.js("""() => { window.showSaveFilePicker = async (o) => { window.__savedAs = o.suggestedName;
            const root = await navigator.storage.getDirectory(); return await root.getFileHandle('settings.json', {create: true}); }; }""")
        self.open_file('t.xlsx', workbook(long_rows(10)))
        self.page.select_option('#pitchSel', '32')
        self.page.click('#settingsBtn')
        self.page.fill('[data-chart="s"][data-n="R2"]', '5')
        self.page.click('#dOk')
        self.page.click('#saveSetBtn')
        self.page.wait_for_function("() => /설정 저장: /.test(document.getElementById('status').textContent)")
        self.assertIn('이 PC 브라우저', self.page.inner_text('#status'))
        self.assertEqual(self.js('() => window.__savedAs'), 'SPC_Chart_설정.json')
        saved = json.loads(base64.b64decode(self.js(GET_FILE, 'settings.json')).decode('utf-8'))
        self.assertEqual((saved['settings']['pitch'], saved['settings']['sigma']['n']), (32, {'R2': 5}))
        self.page.reload()  # closed and opened again: the saved settings are back
        st = self.js('() => window.__spc.settings')
        self.assertEqual((st['pitch'], st['sigma']['n']), (32, {'R2': 5}))
        # a browser that forgot everything: the settings file brings it back
        self.js('() => localStorage.clear()')
        self.page.reload()
        self.assertEqual(self.js('() => window.__spc.settings.pitch'), 16)
        self.page.set_input_files('#setIn', files=[{'name': 'SPC_Chart_설정.json', 'mimeType': 'application/json', 'buffer': json.dumps(saved).encode('utf-8')}])
        self.page.wait_for_function("() => /설정 불러옴: /.test(document.getElementById('status').textContent)")
        st = self.js('() => window.__spc.settings')
        self.assertEqual((st['pitch'], st['sigma']['n']), (32, {'R2': 5}))
        self.assertEqual(json.loads(self.js("() => localStorage.getItem('spc.settings')"))['pitch'], 32)
        self.page.set_input_files('#setIn', files=[{'name': 'x.json', 'mimeType': 'application/json', 'buffer': b'{"a": 1}'}])
        self.page.wait_for_function("() => /설정 파일이 아닙니다/.test(document.getElementById('status').textContent)")

    # issue #25: an exporter's .xlsx that Excel opens but whose zip spells the parts differently
    def test_xlsx_with_other_part_names(self):
        rows = spc_rows()
        def renamed(name, body):  # workbook under another name, found through _rels/.rels
            body = body.replace(b'xl/workbook.xml', b'xl/book.xml').replace(b'workbook.xml.rels', b'book.xml.rels')
            return name.replace('xl/workbook.xml', 'xl/book.xml').replace('workbook.xml.rels', 'book.xml.rels'), body
        def windows(name, body):  # '\\' separators and other letter case
            return name.replace('/', '\\').replace('xl', 'XL').replace('sheet1', 'Sheet1').replace('styles', 'Styles'), body
        for i, rename in enumerate((renamed, windows)):
            data = respell(workbook(rows), rename)
            self.open_file('export%d.xlsx' % i, data)
            self.assertEqual(self.js('() => window.__spc.table().spcForm'), True)
            self.assertEqual(self.results()[0]['n'], len(rows) - 1)
            self.sim('#fixAllBtn')
            saved = bytes(self.js('() => Array.from(window.__spc.copyBytes())'))
            names = zipfile.ZipFile(io.BytesIO(saved)).namelist()
            self.assertEqual(sorted(names), sorted(zipfile.ZipFile(io.BytesIO(data)).namelist()))
            z = zipfile.ZipFile(io.BytesIO(saved))
            sheet = [n for n in names if n.lower().replace('\\', '/').endswith('worksheets/sheet1.xml')][0]
            styles = [n for n in names if n.lower().endswith('styles.xml')][0]
            self.assertIn(b'FFFF0000', z.read(styles))
            self.assertNotEqual(z.read(sheet), zipfile.ZipFile(io.BytesIO(data)).read(sheet))
            if rename is renamed:  # openpyxl follows _rels/.rels too
                ws = openpyxl.load_workbook(io.BytesIO(saved)).active
                self.assertEqual(ws.cell(SPC_OUT + 2, SPC_HEADERS.index('EXT_MV') + 1).fill.fgColor.rgb, 'FFFF0000')

    def test_ods_and_clear_errors(self):
        self.open_file('data.ods', ods(long_rows(30)))
        res = self.results()
        self.assertEqual([r['n'] for r in res], [30, 30])
        self.assertEqual(res[0]['usl'], 10.5)
        # a zip without a workbook names what is inside it; .xlsb says what to do
        msg = self.js("""(b64) => { try { XLSXLite.read('x.xlsx', Uint8Array.from(atob(b64), c => c.charCodeAt(0)).buffer); return ''; }
            catch (e) { return e.message; } }""", base64.b64encode(respell(workbook([[1, 2]]), lambda n, b: (n.replace('workbook', 'wb').replace('xl/', 'q/'), b.replace(b'workbook', b'wb')))).decode())
        self.assertIn('workbook', msg)
        self.assertIn('q/wb.xml', msg)  # names what is inside
        xlsb = io.BytesIO()
        with zipfile.ZipFile(xlsb, 'w') as z:
            z.writestr('[Content_Types].xml', '<Types/>')
            z.writestr('xl/workbook.bin', b'\x83\x01')
        msg = self.js("""(b64) => { try { XLSXLite.read('x.xlsx', Uint8Array.from(atob(b64), c => c.charCodeAt(0)).buffer); return ''; }
            catch (e) { return e.message; } }""", base64.b64encode(xlsb.getvalue()).decode())
        self.assertIn('.xlsb', msg)


if __name__ == '__main__':
    unittest.main()
