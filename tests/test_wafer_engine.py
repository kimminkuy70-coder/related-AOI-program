# -*- coding: utf-8 -*-
"""Wafer Map Converter engine: Excel legend, round trip, discovery, output layout.

    python -m unittest discover -s tests -v     (environment with openpyxl/matplotlib/Pillow)
"""
import importlib
import os
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
APP_DIR = REPO / 'programs' / 'Wafer_Map_Converter' / 'app_files'
os.environ.setdefault('AOI_TOOLS_HOME', tempfile.mkdtemp(prefix='aoi_tools_test_'))

CODES = ['000'] * 12 + ['003', '007', '014', '022', '031', '090', '055']


def write_map(path, wafer, n=12, crlf=True, bom=False):
    rows = []
    c = (n - 1) / 2
    for r in range(n):
        rows.append(['___' if (r - c) ** 2 + (k - c) ** 2 > (n / 2 - .5) ** 2 else CODES[(r * 7 + k * 3) % len(CODES)] for k in range(n)])
    nl = '\r\n' if crlf else '\n'
    head = ['WAFER:' + wafer, 'DEVICE:DEV', 'LOT:LOT1', 'ROWCT:%d' % n, 'COLCT:%d' % n]
    text = nl.join(head + ['RowData:' + ' '.join(r) for r in rows]) + nl
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(('﻿' if bom else '').encode('utf-8') + text.encode('utf-8'))
    return rows


def load(name):
    sys.modules.pop(name, None)
    sys.path.insert(0, str(APP_DIR))
    try:
        return importlib.import_module(name)
    finally:
        sys.path.remove(str(APP_DIR))


class EngineTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        for name in ('engine', 'runtime', 'app'):
            sys.modules.pop(name, None)
        cls.engine = load('engine')

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_legend_next_to_map_uses_map_cell_style(self):
        from openpyxl import load_workbook
        write_map(self.tmp / 'W1.txt', 'W1')
        result = self.engine.txt_to_excel(self.tmp / 'W1.txt', self.tmp / 'W1_Map_Edit.xlsx')
        ws = load_workbook(result['output'])['Map_Edit']
        col = 12 + 3
        self.assertEqual(ws.cell(4, col).value, 'Bin Code')
        legend = {ws.cell(r, col).value: ws.cell(r, col) for r in range(5, 20) if ws.cell(r, col).value}
        for code in ('000', '003', '090', '___', '055'):
            self.assertIn(code, legend)
        self.assertEqual(legend['055'].fill.fgColor.rgb[-6:], 'F4B183')  # unknown code gets the fallback colour
        # a legend cell pasted onto a die must look exactly like a die of that code
        die = next(ws.cell(r, c) for r in range(5, 17) for c in range(2, 14) if ws.cell(r, c).value == '003')
        for attr in ('fillId', 'fontId', 'borderId', 'alignmentId', 'numFmtId'):  # same stored style records
            self.assertEqual(getattr(legend['003']._style, attr), getattr(die._style, attr), attr)
        self.assertEqual(die.number_format, '@')
        self.assertTrue(str(ws.cell(legend['003'].row, col + 4).value).startswith('=SUMPRODUCT'))
        self.assertEqual([Path(i['path']).name for i in result['images']], ['W1_Map_Edit_BinCode_Map.png', 'W1_Map_Edit_BinMeaning_Map.png'])
        self.assertTrue(all(Path(i['path']).exists() for i in result['images']))

    def test_round_trip_is_byte_identical_and_keeps_edits(self):
        from openpyxl import load_workbook
        for crlf, bom in ((True, False), (False, False), (True, True)):
            with self.subTest(crlf=crlf, bom=bom):
                src = self.tmp / ('M%d%d.txt' % (crlf, bom))
                write_map(src, 'M', crlf=crlf, bom=bom)
                xlsx = self.engine.txt_to_excel(src, src.with_name(src.stem + '_Map_Edit.xlsx'))['output']
                back = self.engine.excel_to_txt(xlsx, src.with_name(src.stem + '_Converted.txt'))['output']
                self.assertEqual(Path(back).read_bytes(), src.read_bytes())
        # edits: a pasted code and a number typed without leading zeros
        wb = load_workbook(xlsx)
        ws = wb['Map_Edit']
        ws.cell(10, 7).value = '022'
        ws.cell(10, 8).value = 3
        wb.save(xlsx)
        back = self.engine.excel_to_txt(xlsx, self.tmp / 'edited.txt')['output']
        row = Path(back).read_text(encoding='utf-8-sig').splitlines()[5 + 5].split(':', 1)[1].split()
        self.assertEqual(row[5:7], ['022', '003'])

    def test_dotted_names_keep_their_own_images(self):
        write_map(self.tmp / 'LOT.01.txt', 'A')
        write_map(self.tmp / 'LOT.02.txt', 'B')
        a = self.engine.txt_to_excel(self.tmp / 'LOT.01.txt', self.tmp / 'LOT.01_Map_Edit.xlsx')
        b = self.engine.txt_to_excel(self.tmp / 'LOT.02.txt', self.tmp / 'LOT.02_Map_Edit.xlsx')
        self.assertEqual(Path(a['images'][0]['path']).name, 'LOT.01_Map_Edit_BinCode_Map.png')
        self.assertNotEqual(a['images'][0]['path'], b['images'][0]['path'])

    def test_discovery_per_mode(self):
        write_map(self.tmp / 'a' / 'W1.txt', 'W1')
        self.engine.txt_to_excel(self.tmp / 'a' / 'W1.txt', self.tmp / 'a' / 'W1_Map_Edit.xlsx')
        self.engine.excel_to_txt(self.tmp / 'a' / 'W1_Map_Edit.xlsx', self.tmp / 'a' / 'W1_Map_Edit_Converted.txt')
        (self.tmp / 'a' / '~$W1_Map_Edit.xlsx').write_bytes(b'lock')
        self.assertEqual([p.name for p in self.engine.discover(self.tmp, ('.txt',))], ['W1.txt'])
        self.assertEqual([p.name for p in self.engine.discover(self.tmp, ('.xlsx', '.xlsm'))], ['W1_Map_Edit.xlsx'])


class ApiOutputTest(unittest.TestCase):
    def test_custom_output_keeps_subfolders_and_limits_paths(self):
        import webview
        engine = load('engine')
        load('runtime')
        app = load('app')
        tmp = Path(tempfile.mkdtemp())
        write_map(tmp / 'src' / 'L1' / 'W01.txt', 'A')
        write_map(tmp / 'src' / 'L2' / 'W01.txt', 'B')

        class Window:
            answers = [str(tmp / 'src'), str(tmp / 'out')]

            def create_file_dialog(self, *args, **kwargs):
                return [self.answers.pop(0)]

        api = app.API(webview, engine)
        api._window = Window()
        api.choose_folder('txt')
        self._wait(lambda: api.get_state()['scan'] == 'complete')
        out = api.choose_output()
        self.assertTrue(api.start([0, 1], out)['ok'])
        self._wait(lambda: api.get_state()['run'] != 'running')
        state = api.get_state()
        self.assertEqual(state['errors'], [])
        outputs = sorted(Path(r['output']).relative_to(tmp).as_posix() for r in state['results'])
        self.assertEqual(outputs, ['out/L1/W01_Map_Edit.xlsx', 'out/L2/W01_Map_Edit.xlsx'])
        image = state['results'][0]['images'][0]['path']
        self.assertTrue(api.get_image(image, True).startswith('data:image/jpeg;base64,'))
        self.assertTrue(api.get_image(image).startswith('data:image/png;base64,'))
        self.assertEqual(api.get_image(str(tmp / 'src' / 'L1' / 'W01.txt') + '.nope'), '')
        self.assertFalse(api.open_file(__file__))  # not produced or read by this job

    @staticmethod
    def _wait(done, timeout=120):
        end = time.time() + timeout
        while not done():
            if time.time() > end:
                raise AssertionError('timeout')
            time.sleep(0.1)


if __name__ == '__main__':
    unittest.main()
